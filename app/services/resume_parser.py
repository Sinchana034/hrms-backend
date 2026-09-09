import io
import re
from typing import Any


# Skills that we can detect from resume text.
KNOWN_SKILLS = [
    "python",
    "java",
    "javascript",
    "typescript",
    "c",
    "c++",
    "c#",
    "html",
    "css",
    "react",
    "reactjs",
    "angular",
    "vue",
    "node",
    "node.js",
    "express",
    "fastapi",
    "flask",
    "django",
    "sql",
    "mysql",
    "postgresql",
    "postgres",
    "mongodb",
    "git",
    "github",
    "docker",
    "kubernetes",
    "aws",
    "azure",
    "gcp",
    "machine learning",
    "deep learning",
    "tensorflow",
    "pytorch",
    "pandas",
    "numpy",
    "scikit-learn",
    "rest api",
    "rest",
    "api",
    "spring boot",
    "spring",
    "bootstrap",
    "tailwind",
]


def extract_pdf_text(file_bytes: bytes) -> str:
    """
    Extract text from a PDF resume.
    """
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(file_bytes))

    pages = []

    for page in reader.pages:
        text = page.extract_text()

        if text:
            pages.append(text)

    return "\n".join(pages)


def extract_docx_text(file_bytes: bytes) -> str:
    """
    Extract text from a DOCX resume.
    """
    from docx import Document

    document = Document(io.BytesIO(file_bytes))

    paragraphs = []

    for paragraph in document.paragraphs:
        if paragraph.text:
            paragraphs.append(paragraph.text)

    # Also read tables because many resumes use tables.
    for table in document.tables:
        for row in table.rows:
            for cell in row.cells:
                if cell.text:
                    paragraphs.append(cell.text)

    return "\n".join(paragraphs)


def extract_doc_text(file_bytes: bytes) -> str:
    """
    Extract text from an old .doc file.

    This requires the `antiword` command to be installed.
    """
    import subprocess

    try:
        result = subprocess.run(
            ["antiword", "-"],
            input=file_bytes,
            stdout=subprocess.PIPE,
            stderr=subprocess.PIPE,
            check=True,
        )

        return result.stdout.decode("utf-8", errors="ignore")

    except FileNotFoundError:
        raise RuntimeError(
            "DOC resume parsing requires antiword. "
            "Install it with: brew install antiword"
        )

    except subprocess.CalledProcessError as exc:
        raise RuntimeError(
            f"Failed to parse DOC resume: "
            f"{exc.stderr.decode('utf-8', errors='ignore')}"
        )


def extract_text(file_bytes: bytes, storage_path: str) -> str:
    """
    Determine file type from the storage path and extract text.
    """

    extension = storage_path.lower().split(".")[-1]

    if extension == "pdf":
        return extract_pdf_text(file_bytes)

    if extension == "docx":
        return extract_docx_text(file_bytes)

    if extension == "doc":
        return extract_doc_text(file_bytes)

    raise RuntimeError(
        f"Unsupported resume format: .{extension}"
    )


def normalize_text(text: str) -> str:
    """
    Normalize resume text for easier matching.
    """
    text = text.lower()

    # Replace punctuation with spaces, but preserve + and #.
    text = re.sub(r"[(),;|:/]", " ", text)

    # Normalize whitespace.
    text = re.sub(r"\s+", " ", text)

    return text.strip()


def extract_skills(text: str) -> list[str]:
    """
    Detect known technical skills in resume text.
    """

    normalized = normalize_text(text)

    found = []

    # Longer/multi-word skills first.
    sorted_skills = sorted(
        KNOWN_SKILLS,
        key=len,
        reverse=True,
    )

    for skill in sorted_skills:

        skill_normalized = normalize_text(skill)

        pattern = r"(?<![a-z0-9])" + re.escape(skill_normalized) + r"(?![a-z0-9])"

        if re.search(pattern, normalized):
            canonical_skill = skill.lower()

            # Normalize aliases.
            if canonical_skill == "reactjs":
                canonical_skill = "react"

            elif canonical_skill == "node.js":
                canonical_skill = "node"

            elif canonical_skill == "postgres":
                canonical_skill = "postgresql"

            if canonical_skill not in found:
                found.append(canonical_skill)

    return sorted(found)


def extract_section(
    text: str,
    section_names: list[str],
) -> str:
    """
    Extract text belonging to a resume section.

    Example:
        SKILLS
        Python
        FastAPI
        Git

        EDUCATION
        B.E Computer Science
    """

    lines = [
        line.strip()
        for line in text.splitlines()
        if line.strip()
    ]

    if not lines:
        return ""

    normalized_names = {
        name.lower()
        for name in section_names
    }

    start_index = None

    for index, line in enumerate(lines):

        cleaned = re.sub(
            r"[^a-zA-Z ]",
            "",
            line,
        ).strip().lower()

        if cleaned in normalized_names:
            start_index = index + 1
            break

    if start_index is None:
        return ""

    section_lines = []

    possible_next_sections = {
        "skills",
        "technical skills",
        "education",
        "experience",
        "work experience",
        "professional experience",
        "projects",
        "project",
        "certifications",
        "achievements",
        "summary",
        "objective",
        "contact",
    }

    for line in lines[start_index:]:

        cleaned = re.sub(
            r"[^a-zA-Z ]",
            "",
            line,
        ).strip().lower()

        if cleaned in possible_next_sections:
            break

        section_lines.append(line)

    return "\n".join(section_lines)


def extract_education(text: str) -> list[dict[str, Any]]:
    """
    Basic education extraction.
    """

    section = extract_section(
        text,
        [
            "education",
            "academic background",
            "academic qualifications",
        ],
    )

    if not section:
        return []

    results = []

    for line in section.splitlines():

        line = line.strip()

        if not line:
            continue

        results.append(
            {
                "details": line
            }
        )

    return results[:10]


def extract_experience(text: str) -> list[dict[str, Any]]:
    """
    Basic work-experience extraction.
    """

    section = extract_section(
        text,
        [
            "experience",
            "work experience",
            "professional experience",
            "employment",
        ],
    )

    if not section:
        return []

    results = []

    for line in section.splitlines():

        line = line.strip()

        if not line:
            continue

        results.append(
            {
                "details": line
            }
        )

    return results[:20]


def extract_projects(text: str) -> list[dict[str, Any]]:
    """
    Basic project extraction.
    """

    section = extract_section(
        text,
        [
            "projects",
            "project",
            "academic projects",
            "personal projects",
        ],
    )

    if not section:
        return []

    results = []

    for line in section.splitlines():

        line = line.strip()

        if not line:
            continue

        results.append(
            {
                "details": line
            }
        )

    return results[:20]


def parse_resume(
    file_bytes: bytes,
    storage_path: str,
) -> dict[str, Any]:
    """
    Complete resume parsing pipeline.
    """

    text = extract_text(
        file_bytes,
        storage_path,
    )

    if not text.strip():
        raise RuntimeError(
            "Could not extract any text from the resume. "
            "If this is a scanned/image-only PDF, OCR is required."
        )

    skills = extract_skills(text)

    education = extract_education(text)

    experience = extract_experience(text)

    projects = extract_projects(text)

    return {
        "skills": skills,
        "education": education,
        "experience": experience,
        "projects": projects,
        "resume_text": text,
    }